import os
import json
import uuid
import datetime
import io
from typing import List, Dict, Optional
import pypdf
import docx

DOCUMENTS_FILE_PATH = "data/documents.json"
UPLOAD_DIR = ".tmp/uploads"

def get_documents() -> List[Dict]:
    """Retrieves all indexed documents from storage."""
    if not os.path.exists(DOCUMENTS_FILE_PATH):
        return []
    try:
        with open(DOCUMENTS_FILE_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    except (json.JSONDecodeError, OSError):
        return []

def save_document_record(doc: Dict) -> bool:
    """Appends or updates a document record in data/documents.json."""
    docs = get_documents()
    # Replace if ID exists, otherwise append
    docs = [d for d in docs if d.get("id") != doc.get("id")]
    docs.append(doc)

    os.makedirs(os.path.dirname(DOCUMENTS_FILE_PATH), exist_ok=True)
    try:
        with open(DOCUMENTS_FILE_PATH, "w", encoding="utf-8") as f:
            json.dump(docs, f, indent=4)
        return True
    except Exception as e:
        print(f"Error saving documents list: {e}")
        return False

def delete_document(doc_id: str) -> bool:
    """Removes a document by ID."""
    docs = get_documents()
    filtered = [d for d in docs if d.get("id") != doc_id]
    if len(filtered) == len(docs):
        return False
    os.makedirs(os.path.dirname(DOCUMENTS_FILE_PATH), exist_ok=True)
    with open(DOCUMENTS_FILE_PATH, "w", encoding="utf-8") as f:
        json.dump(filtered, f, indent=4)
    return True

def extract_text_from_bytes(filename: str, file_bytes: bytes) -> str:
    """Extracts raw text from PDF, DOCX, or TXT bytes."""
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == ".pdf":
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        text_parts = []
        for page_num, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(page_text.strip())
        return "\n\n".join(text_parts)
    
    elif ext in [".docx", ".doc"]:
        doc = docx.Document(io.BytesIO(file_bytes))
        text_parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
        return "\n\n".join(text_parts)
    
    elif ext in [".txt", ".md", ".csv"]:
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1", errors="replace")
    
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Please upload a PDF, DOCX, or TXT file.")

def process_and_store_document(filename: str, file_bytes: bytes) -> Dict:
    """
    Parses document content, saves upload to .tmp/uploads, and indexes in data/documents.json.
    """
    extracted_text = extract_text_from_bytes(filename, file_bytes)
    
    if not extracted_text.strip():
        raise ValueError("The uploaded document appears to be empty or unreadable.")

    # Persist file copy in .tmp/uploads
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    doc_id = str(uuid.uuid4())[:8]
    sanitized_filename = f"{doc_id}_{filename}"
    upload_path = os.path.join(UPLOAD_DIR, sanitized_filename)
    
    with open(upload_path, "wb") as f:
        f.write(file_bytes)

    word_count = len(extracted_text.split())
    char_count = len(extracted_text)

    doc_record = {
        "id": doc_id,
        "filename": filename,
        "upload_path": upload_path,
        "uploaded_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "word_count": word_count,
        "char_count": char_count,
        "content": extracted_text
    }

    save_document_record(doc_record)
    return doc_record

def get_combined_document_text(max_chars: int = 15000) -> str:
    """Returns aggregated text excerpt from all indexed documents."""
    docs = get_documents()
    if not docs:
        return ""
    
    aggregated = []
    current_len = 0
    for doc in docs:
        header = f"=== DOCUMENT: {doc.get('filename')} ==="
        body = doc.get('content', '')
        snippet = f"{header}\n{body}\n"
        if current_len + len(snippet) > max_chars:
            remaining = max(0, max_chars - current_len - len(header) - 20)
            if remaining > 100:
                aggregated.append(f"{header}\n{body[:remaining]}... [TRUNCATED]")
            break
        aggregated.append(snippet)
        current_len += len(snippet)

    return "\n\n".join(aggregated)
